"""
Author:  Kevin Thomas, Director or Admissions/Registrar
Date:    December, 2023
Purpose: This program will store and print student diplomas  

Directions:  
    
my need to pip install sevaral libraries including PIP INSTALL PACKAGING

"""

"""import os
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd
from pandas import ExcelWriter
from pandas.io.parsers import read_csv
from xlsxwriter import *"""
import openpyxl as opxl
import tkinter as tk
from tkinter import *
import customtkinter
import sys
from tkinter import ttk
#import csv
#from customtkinter import *


class App(customtkinter.CTk):
    def __init__(self):
        super().__init__()
        customtkinter.set_default_color_theme("green")  # Themes: "blue" (standard), "green", "dark-blue"
        customtkinter.set_appearance_mode("dark")
        
        self.title('NSCC Credential Generator')
        self.propagate(YES)
        #self.grid_columnconfigure(0, weight=1)
        #self.grid_rowconfigure(0, weight=1)

        #Create a Frame Named Students and Awards and Place these buttons in the frame
        self.studentAndAwardsFrame=customtkinter.CTkFrame(self)
        self.studentAndAwardsFrame.grid(row=0, column=0, padx=10, pady=(0, 0), sticky="nsew")
        
        #Create new student button
        self.newStudent_button = customtkinter.CTkButton(self.studentAndAwardsFrame, text="Create New Student", command=self.createNewStudent)
        self.newStudent_button.grid(row=1, column=1, padx=10, pady=10)

        #Create import students and awards button
        self.importStudentsAndAwardsButton = customtkinter.CTkButton(self.studentAndAwardsFrame, text="Import Students and Awards", command=self.importStudentsAndAwards)
        self.importStudentsAndAwardsButton.grid(row=1, column=2, padx=10, pady=10)

        #Display excel spreadsheet
        treeFrame = ttk.Frame(self)
        treeFrame.grid(row=2, column=0, padx=5, pady=(0, 0), sticky="nsew")

        treeScroll = ttk.Scrollbar(treeFrame)
        treeScroll.pack(side="right", fill='y')

        self.cols = ("student_id","term","grad_date","first_name","middle_name","last_name","full_name","degree_1","major_1","honor_1","concentration","other")

        self.treeView = ttk.Treeview(treeFrame, show='headings', yscrollcommand=treeScroll.set, columns=self.cols, height=30)

        self.treeView.column("student_id", width=100)
        self.treeView.column("term", width=75)
        self.treeView.column("grad_date", width=100)
        self.treeView.column("first_name", width=200)
        self.treeView.column("middle_name", width=200)
        self.treeView.column("last_name", width=200)
        self.treeView.column("full_name", width=200)
        self.treeView.column("degree_1", width=200)
        self.treeView.column("major_1", width=225)
        self.treeView.column("honor_1", width=200)
        self.treeView.column("concentration", width=200)
        self.treeView.column("other", width=75)

        self.treeView.pack()
        treeScroll.config(command=self.treeView.yview)

        self.loadData()

        #Create button to exit the program
        self.endButton = customtkinter.CTkButton(self, text="Close", command=self.closeProgram)
        self.endButton.grid(row=88, column=88, padx=10, pady=(10,0), sticky='ew')

    def createNewStudent(self):
        enterGradRecord=customtkinter.CTk()
        enterGradRecord.title("Enter a New Graduate")
        enterGradRecord.pack_propagate(YES)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)


        def save():
            #Get Data from form
            id=student_id_entry.get()
            gradTerm=term_entry.get()
            gradDate=grad_date_entry.get()
            firstName=first_name_entry.get()
            middleName=middle_name_entry.get()
            lastName=last_name_entry.get()
            fullName=full_name_entry.get()
            degree=degree_1_entry.get()
            major=major_1_entry.get()
            honor=honor_1_entry.get()
            concentration = concentration_entry.get()
            other=other_entry.get()

            #Insert Row into the Excel File
            file="Excel for Python Credential Generator.xlsx"
            workbook = opxl.load_workbook(file)
            sheet = workbook.active

            row = [id,gradTerm,gradDate,firstName,middleName,lastName,fullName,degree,major,concentration, honor,other]

            sheet.append(row)
            workbook.save(file)
            print("Graduate Created")

            #Insert row into 
            self.treeView.insert('', tk.END, values=row)

            #Reset Form
            student_id_entry.delete(0,"end")
            student_id_entry.insert(0,"Enter Student ID")

            term_entry.delete(0, "end")
            term_entry.insert(0,"Enter Grad Term")

            grad_date_entry.delete(0, "end")
            grad_date_entry.insert(0,"Enter Grad Date")

            first_name_entry.delete(0, "end")
            first_name_entry.insert(0,"Enter First Name")

            middle_name_entry.delete(0,'end')
            middle_name_entry.insert(0, 'Enter Middle Name')

            last_name_entry.delete(0,"end")
            last_name_entry.insert(0,"Enter Last Name")

            full_name_entry.delete(0,'end')
            full_name_entry.insert(0, "Enter Full Name")

            degree_1_entry.delete(0,'end')
            degree_1_entry.insert(0,'Enter Degree')

            major_1_entry.delete(0,'end')
            major_1_entry.insert(0,'Enter Major')

            concentration_entry.delete(0,'end')
            concentration_entry.insert(0,'Enter concentration')

            honor_1_entry.delete(0,'end')
            honor_1_entry.insert(0,'Enter Honors')

            other_entry.delete(0,'end')
            other_entry.insert(0, 'Enter Other')

        def close():
            enterGradRecord.destroy()

        #student information frame
        studentInfoFrame = customtkinter.CTkFrame(enterGradRecord)
        studentInfoFrame.grid(row=0, column=0)
        testLabel=Label(studentInfoFrame, text="Enter Graduate Information", padx=10, pady=10)
        testLabel.grid(row=0, column=0, columnspan=5)

        student_id_entry=ttk.Entry(studentInfoFrame)
        student_id_entry.insert(0,"Enter Student ID")
        student_id_entry.bind("<FocusIn>", lambda e: student_id_entry.delete(0,END))
        student_id_entry.grid(row=1, column=0, padx=10, pady=10, sticky='ew')

        term_entry=ttk.Entry(studentInfoFrame)
        term_entry.insert(0,"Enter Grad Term")
        term_entry.bind("<FocusIn>", lambda e: term_entry.delete(0,END))
        term_entry.grid(row=2, column=0,  padx=10, pady=10, sticky='ew')

        grad_date_entry=ttk.Entry(studentInfoFrame)
        grad_date_entry.insert(0,"Enter Grad Date")
        grad_date_entry.bind("<FocusIn>", lambda e: grad_date_entry.delete(0,END))
        grad_date_entry.grid(row=3, column=0, padx=10, pady=10,  sticky='ew')

        first_name_entry=ttk.Entry(studentInfoFrame)
        first_name_entry.insert(0,"Enter First Name")
        first_name_entry.bind("<FocusIn>", lambda e: first_name_entry.delete(0,END))
        first_name_entry.grid(row=4, column=0, padx=10, pady=10,  sticky='ew')

        middle_name_entry=ttk.Entry(studentInfoFrame)
        middle_name_entry.insert(0,"Enter Middle Name")
        middle_name_entry.bind("<FocusIn>", lambda e: middle_name_entry.delete(0,END))
        middle_name_entry.grid(row=5, column=0, padx=10, pady=10,  sticky='ew')

        last_name_entry=ttk.Entry(studentInfoFrame)
        last_name_entry.insert(0,"Enter Last Name")
        last_name_entry.bind("<FocusIn>", lambda e: last_name_entry.delete(0,END))
        last_name_entry.grid(row=1, column=1, padx=10, pady=10,  sticky='ew') 

        full_name_entry=ttk.Entry(studentInfoFrame)
        full_name_entry.insert(0,"Enter Full Name")
        full_name_entry.bind("<FocusIn>", lambda e: full_name_entry.delete(0,END))
        full_name_entry.grid(row=2, column=1, padx=10, pady=10,  sticky='ew')

        degree_1_entry=ttk.Entry(studentInfoFrame)
        degree_1_entry.insert(0,"Enter Degree")
        degree_1_entry.bind("<FocusIn>", lambda e: degree_1_entry.delete(0,END))
        degree_1_entry.grid(row=3, column=1, padx=10, pady=10,  sticky='ew')  

        major_1_entry=ttk.Entry(studentInfoFrame)
        major_1_entry.insert(0,"Enter Major")
        major_1_entry.bind("<FocusIn>", lambda e: major_1_entry.delete(0,END))
        major_1_entry.grid(row=4, column=1, padx=10, pady=10,  sticky='ew') 

        concentration_entry=ttk.Entry(studentInfoFrame)
        concentration_entry.insert(0,"Enter Concentration Code")
        concentration_entry.bind("<FocusIn>", lambda e: concentration_entry.delete(0,END))
        concentration_entry.grid(row=5, column=1, padx=10, pady=10,  sticky='ew') 

        honor_1_entry=ttk.Entry(studentInfoFrame)
        honor_1_entry.insert(0,"Enter Honor Code")
        honor_1_entry.bind("<FocusIn>", lambda e: honor_1_entry.delete(0,END))
        honor_1_entry.grid(row=5, column=1, padx=10, pady=10,  sticky='ew') 

        other_entry=ttk.Entry(studentInfoFrame)
        other_entry.insert(0,"Enter Other Information")
        other_entry.bind("<FocusIn>", lambda e: other_entry.delete(0,END))
        other_entry.grid(row=6, column=1, padx=10, pady=10,  sticky='ew')   


        #controls frame
        studentControlFrame= customtkinter.CTkFrame(enterGradRecord)
        studentControlFrame.grid(row=88, column=88)

        saveButton = customtkinter.CTkButton(studentControlFrame, text="Save", command=save)
        saveButton.grid(row=0, column=0, padx=10, pady=20)

        backButton = customtkinter.CTkButton(studentControlFrame, text="Back", command=close)
        backButton.grid(row=0, column=1, padx=10, pady=20)

        enterGradRecord.mainloop()

    def importStudentsAndAwards(self):
        #Load file of original data
        dbFile='Excel for Python Credential Generator.xlsx'
        dbFileWorkBook=opxl.load_workbook(dbFile)
        dbSheet = dbFileWorkBook.active

        #Load file of new data
        importFile='Test Credential Import File.xlsx'
        importWorkBook=opxl.load_workbook(importFile)
        sheet=importWorkBook.active

        row_count = sheet.max_row
        col_count = sheet.max_column

        list_of_lists = []

        for row in range(2, row_count+1):
            list=[]
            for column in range(1, col_count+1):
                val=sheet.cell(row=row, column=column).value
                list.append(val)
            list_of_lists.append(list)

        for row in list_of_lists:
            dbSheet.append(row)
        
        dbFileWorkBook.save('Excel for Python Credential Generator.xlsx')
        print("The new file has been added to the database")

    def createMergedMultipleCampsFile(self):
        print("Button Pressed")
    
    def closeProgram(self):
        print("Exiting the program...")
        sys.exit(0)
    
    def loadData(self):
        file="Excel for Python Credential Generator.xlsx"
        workbook = opxl.load_workbook(file)
        sheet = workbook.active

        list_values= list(sheet.values)

        for col_name in list_values[0]:
            self.treeView.heading(col_name, text=col_name)
        
        for row_data in list_values[1:]:
            self.treeView.insert('', tk.END, values=row_data)

#Criteria to search by drop down list (Term, A#, LastName)

#Data entry box to enter text that matches Criteria (202310, A00495006, Thomas)

app = App()
app.mainloop()
